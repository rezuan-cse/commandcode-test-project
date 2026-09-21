"""Workbook parsing. Reads the client's Excel file into plain dictionaries.

Kept separate from loading so the parsing rules (column positions, code
normalisation, Excel serial dates) are testable on their own.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXCEL_EPOCH = dt.date(1899, 12, 30)

CATEGORY_MAP = {
    "raw material": "Raw Material",
    "wip": "WIP",
    "finished good": "Finished Good",
    "packaging": "Packaging",
    "trading stock": "Trading Stock",
    "trading": "Trading Stock",
    "imported goods": "Imported Goods",
}

MOVEMENT_MAP = {
    "opening stock": "Opening Stock",
    "purchase-in": "Purchase-In",
    "production-in": "Production-In",
    "production-out": "Production-Out",
    "sale-out": "Sale-Out",
    "adjustment": "Adjustment",
}


@dataclass
class WorkbookData:
    """Everything extracted from the workbook, before it touches the database."""

    accounts: list[dict[str, Any]] = field(default_factory=list)
    opening_balances: list[dict[str, Any]] = field(default_factory=list)
    cutover_date: dt.date | None = None
    journal_entries: list[dict[str, Any]] = field(default_factory=list)
    items: list[dict[str, Any]] = field(default_factory=list)
    bom: list[dict[str, Any]] = field(default_factory=list)
    inventory: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def is_valid_code(text: str) -> bool:
    """True when text looks like a real code rather than a header or label.

    Sheet headers and total rows are written in Bengali and contain spaces and
    parentheses, so rejecting those characters is enough to skip them.
    """
    if not text:
        return False
    return all(char.isalnum() or char in "-_." for char in text)


def normalise_code(value: Any) -> str:
    """Turn ``1000.0`` into ``1000``, returning ``""`` for headers and blanks."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "none":
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    return text if is_valid_code(text) else ""


def to_date(value: Any) -> dt.date | None:
    """Convert an Excel serial number or datetime into a :class:`date`."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return EXCEL_EPOCH + dt.timedelta(days=int(float(value)))
    except (TypeError, ValueError):
        return None


def to_number(value: Any) -> str:
    """Return a numeric cell as a clean string, or ``0`` when blank."""
    if value is None or value == "":
        return "0"
    return str(value)


def _parse_accounts(ws: Any, data: WorkbookData) -> None:
    """Sheet 2: account code, name, type, segment, normal balance."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = normalise_code(row[0] if row else None)
        if not code:
            continue
        data.accounts.append(
            {
                "code": code,
                "name_en": str(row[1]).strip() if row[1] else code,
                "account_type": str(row[2]).strip() if row[2] else "Asset",
                "segment": str(row[3]).strip() if row[3] else "Shared",
                "normal_balance": str(row[4]).strip() if row[4] else "Dr",
            }
        )


def _parse_opening_balances(ws: Any, data: WorkbookData) -> None:
    """Sheet 3: cutover date, then account rows with debit and credit."""
    cutover_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if cutover_row and len(cutover_row) > 1:
        data.cutover_date = to_date(cutover_row[1])

    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row:
            continue
        code = normalise_code(row[0])
        if not code:
            continue
        data.opening_balances.append(
            {
                "account_code": code,
                "debit": to_number(row[2]),
                "credit": to_number(row[3]),
            }
        )


def _parse_journal_entries(ws: Any, data: WorkbookData) -> None:
    """Sheet 4: one row per journal line, grouped by voucher number."""
    grouped: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        voucher = normalise_code(row[2])
        code = normalise_code(row[3])
        if not voucher or not code:
            continue
        amount_debit = to_number(row[6])
        amount_credit = to_number(row[7])
        if amount_debit == "0" and amount_credit == "0":
            continue
        entry = grouped.setdefault(
            voucher,
            {
                "voucher_no": voucher,
                "entry_date": to_date(row[1]),
                "posted_by": str(row[9]).strip() if row[9] else "Store",
                "narration": str(row[8]).strip() if row[8] else voucher,
                "lines": [],
            },
        )
        entry["lines"].append(
            {
                "account_code": code,
                "segment": str(row[5]).strip() if row[5] else "Shared",
                "debit": amount_debit,
                "credit": amount_credit,
                "narration": str(row[8]).strip() if row[8] else None,
            }
        )
    data.journal_entries = list(grouped.values())


def _parse_items(ws: Any, data: WorkbookData) -> None:
    """Sheet 5: item code, name, category, segment, unit."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = normalise_code(row[0] if row else None)
        if not code:
            continue
        raw_category = str(row[2]).strip().lower() if row[2] else ""
        data.items.append(
            {
                "code": code,
                "name": str(row[1]).strip() if row[1] else code,
                "category": CATEGORY_MAP.get(raw_category, "Raw Material"),
                "segment": str(row[3]).strip() if row[3] else "Manufacturing",
                "uom": str(row[4]).strip() if row[4] else "pcs",
            }
        )


def _parse_bom(ws: Any, data: WorkbookData) -> None:
    """Sheet 6: parent, stage, component, quantity per unit."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        parent = normalise_code(row[0])
        component = normalise_code(row[3])
        if not parent or not component:
            continue
        data.bom.append(
            {
                "parent_code": parent,
                "component_code": component,
                "qty_per_unit": to_number(row[5]),
                "stage": int(row[2]) if row[2] else 1,
            }
        )


def _parse_inventory(ws: Any, data: WorkbookData) -> None:
    """Sheet 8: movements that actually carry a quantity.

    Placeholder rows with no in/out quantity are skipped; the workbook contains
    many of them where the client prepared a layout but never entered data.
    """
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        item_code = normalise_code(row[1])
        if not item_code:
            continue
        in_qty = to_number(row[5])
        out_qty = to_number(row[7])
        if in_qty == "0" and out_qty == "0":
            continue
        raw_type = str(row[3]).strip().lower() if row[3] else ""
        data.inventory.append(
            {
                "movement_date": to_date(row[0]) or data.cutover_date,
                "item_code": item_code,
                "movement_type": MOVEMENT_MAP.get(raw_type, "Adjustment"),
                "reference": normalise_code(row[4]) or None,
                "in_qty": in_qty,
                "in_value": to_number(row[6]),
                "out_qty": out_qty,
                "out_value": to_number(row[8]),
                "suggested_out_value": to_number(row[9]),
                "workbook_avg_cost": to_number(row[12]),
            }
        )


def load_workbook_data(path: str | Path) -> WorkbookData:
    """Parse every relevant sheet of the client workbook."""
    data = WorkbookData()
    wb = load_workbook(filename=str(path), data_only=True)

    _parse_accounts(wb["Chart of Accounts"], data)
    _parse_opening_balances(wb["Opening Balances"], data)
    _parse_journal_entries(wb["Journal Entries"], data)
    _parse_items(wb["Item Master"], data)
    _parse_bom(wb["BOM Master"], data)
    _parse_inventory(wb["Inventory Ledger"], data)

    if not data.accounts:
        data.warnings.append("No accounts found in the workbook")
    if not data.cutover_date:
        data.warnings.append("No cutover date found; defaulting to today")
    return data
